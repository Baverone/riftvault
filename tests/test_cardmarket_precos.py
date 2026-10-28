"""Os preços do Cardmarket, todos os dias, e o Excel gerado deles (2026-10-25).

André: *"preciso que estejas constantemente a atulizar o Price Trend do market,
principalmente de Radiance [...] podes criar Excel para servir de base de dados
e ires atualizando os precos"*.

O que se fixa:
  1. a LIGAÇÃO: cada impressão do catálogo liga-se ao produto do Cardmarket pelo
     `cmid` da dotgg (pelo id exacto ou por alias); com duas linhas da dotgg na
     mesma impressão fica a que tem `cmid`; sem a dotgg usa-se a ligação gravada;
  2. o que se GUARDA é o guia oficial do Cardmarket — o `trend` a sério — e NÃO
     o `cmPrice` da dotgg, que é a média de 1 dia (medido a 25/10); um zero do
     guia é `NULL`; o dia é o do guia;
  3. uma linha por impressão e por dia: o mesmo dia substitui, dias diferentes
     acumulam, e nada toca na Coleção nem no `price_history` do CardTrader;
  4. o preço de referência: numa carta que só existe em foil é o trend do foil;
     o valor não inventa um preço que falta; a variação só com medição dentro
     da janela;
  5. o EXCEL: as cinco folhas, a data e a hora lá dentro, a nota «gerado, nunca
     editado à mão», a RAD à parte, «O que tenho» só com o que ele tem, o lado
     normal em branco numa carta só-foil, o ficheiro de nome fixo — e gerá-lo
     não escreve na base.

Contra pastas temporárias, sem rede (as duas fontes injectam-se).
"""

from __future__ import annotations

import contextlib
import hashlib
import io
import os
import sys
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from tests.fixture import REPO, SEM_CONFIG, Vault, ficheiro_temporario  # noqa: E402


def guia(dia="2026-10-25", **por_produto):
    """Um price guide de brincar com a forma do oficial."""
    linhas = []
    for pid, campos in por_produto.items():
        linhas.append({"idProduct": int(pid.lstrip("p")), "idCategory": 1655, **campos})
    return {"version": 1, "createdAt": f"{dia}T02:47:35+0200", "priceGuides": linhas}


def dg(gid, cmid=None, has_normal="1", cm_price="9.99", **extra):
    return {"id": gid, "cmid": str(cmid) if cmid else None,
            "cmurl": f"https://www.cardmarket.com/en/Riftbound/Products/Singles/X/{gid}" if cmid else None,
            "cmPrice": cm_price, "cmFoilPrice": "8.88", "hasNormal": has_normal,
            "hasFoil": "1", "price": "0.50", "foilPrice": "1.50", **extra}


#: o Trend e as médias — números que nunca se confundem com o `cmPrice` (9,99).
NORMAL = {"avg": 0.4, "low": 0.02, "trend": 0.42, "avg1": 0.39, "avg7": 0.41,
          "avg30": 0.39, "avg-foil": 2.4, "low-foil": 0.2, "trend-foil": 2.12,
          "avg1-foil": 1.99, "avg7-foil": 2.01, "avg30-foil": 2.37}
SO_FOIL = {"avg": None, "low": 39.41, "trend": 0, "avg1": 95, "avg7": 92, "avg30": 92,
           "avg-foil": 69.86, "low-foil": 39.41, "trend-foil": 60.61,
           "avg1-foil": 60, "avg7-foil": 61.1, "avg30-foil": 73.94}


class Base(unittest.TestCase):

    def setUp(self):
        os.environ["RIFTVAULT_CONFIG"] = SEM_CONFIG
        self.v = Vault()
        from riftvault import cardmarket_precos, collection
        self.cmp = cardmarket_precos
        con = self.v.connect()
        self.v.add_printing(con, "rad-001", "RAD", 1, "Amateur Demolitionist")
        self.v.add_printing(con, "rad-113", "RAD", 113, "Angler Wyrm", rarity="epic")
        self.v.add_printing(con, "tst-001-100", "TST", 1, "Defy", size=100)
        self.v.add_printing(con, "tst-002-100", "TST", 2, "Sem Produto", size=100)
        self.v.rebuild(con)
        collection.adjust(con, "rad-001", 3, source="test")
        collection.adjust(con, "rad-113", 2, source="test")
        con.execute("UPDATE copies SET qty_foil = 1 WHERE printing_id = 'rad-001'")
        con.close()
        self.guia = guia(p884113=NORMAL, p898148=SO_FOIL, p1=NORMAL)
        self.dotgg = [dg("RAD-001", 884113), dg("RAD-113", 898148, has_normal="0"),
                      dg("TST-001", 1), dg("TST-001/100", None), dg("TST-002", None),
                      dg("XXX-999", 5)]

    def tearDown(self):
        self.v.close()
        os.environ.pop("RIFTVAULT_CONFIG", None)

    def sync(self, g=None, d="__", **kw):
        con = self.v.connect()
        try:
            return self.cmp.sincronizar(con, guia=g or self.guia,
                                        dotgg=self.dotgg if d == "__" else d,
                                        log=lambda *_: None, **kw)
        finally:
            con.close()

    def linhas(self):
        con = self.v.connect()
        try:
            return {pid: s for pid, s in self.cmp.series(con).items()}
        finally:
            con.close()


class TestALigacao(Base):

    def test_liga_pelo_id_exacto_e_por_alias(self):
        r = self.sync()
        ids = {ln["printing_id"] for ln in r["linhas"]}
        self.assertEqual(ids, {"rad-001", "rad-113", "tst-001-100"})

    def test_a_duplicada_fica_com_a_que_tem_cmid(self):
        # «TST-001» e «TST-001/100» caem na mesma impressão; só a primeira tem cmid
        self.dotgg = [dg("TST-001/100", None), dg("TST-001", 1)]
        r = self.sync()
        self.assertEqual([ln["cm_product_id"] for ln in r["linhas"]], [1])

    def test_o_que_nao_tem_impressao_ou_produto_conta_se(self):
        r = self.sync()
        self.assertEqual(r["sem_impressao"], 1)        # XXX-999
        self.assertEqual(r["sem_produto"], 1)          # TST-002 sem cmid

    def test_sem_a_dotgg_usa_a_ligacao_gravada(self):
        self.sync()
        r = self.sync(g=guia("2026-10-26", p884113=NORMAL, p898148=SO_FOIL, p1=NORMAL),
                      d=None, sem_dotgg=True)
        self.assertEqual(r["origem_ligacao"], "ligacao-anterior")
        self.assertEqual(r["gravadas"], 3)
        self.assertEqual(len(self.linhas()["rad-001"]), 2)


class TestOQueSeGuarda(Base):

    def test_guarda_o_trend_do_guia_e_nao_o_cmprice_da_dotgg(self):
        self.sync()
        ln = self.linhas()["rad-001"][-1]
        self.assertEqual(ln["trend_cents"], 42)
        self.assertEqual(ln["avg1_cents"], 39)
        self.assertEqual(ln["avg30_cents"], 39)
        self.assertEqual(ln["trend_foil_cents"], 212)
        self.assertNotIn(999, [v for k, v in ln.items() if k.endswith("_cents")])

    def test_mudar_o_cmprice_da_dotgg_nao_muda_nada_do_que_se_guarda(self):
        self.sync()
        antes = self.linhas()
        self.dotgg = [dict(d, cmPrice="123.45", cmFoilPrice="67.89") for d in self.dotgg]
        self.sync()
        self.assertEqual(antes, self.linhas())

    def test_um_zero_do_guia_e_null(self):
        self.sync()
        self.assertIsNone(self.linhas()["rad-113"][-1]["trend_cents"])

    def test_o_dia_e_o_do_guia(self):
        r = self.sync(g=guia("2026-10-20", p884113=NORMAL))
        self.assertEqual(r["dia"], "2026-10-20")
        self.assertEqual(self.linhas()["rad-001"][-1]["day"], "2026-10-20")

    def test_o_outro_mercado_fica_em_centimos_de_dolar(self):
        self.sync()
        ln = self.linhas()["rad-001"][-1]
        self.assertEqual((ln["outro_usd_cents"], ln["outro_usd_foil_cents"]), (50, 150))


class TestUmaLinhaPorDia(Base):

    def test_o_mesmo_dia_substitui_e_dias_diferentes_acumulam(self):
        self.sync()
        self.sync()
        self.assertEqual(len(self.linhas()["rad-001"]), 1)
        self.sync(g=guia("2026-10-26", p884113=NORMAL))
        self.assertEqual([ln["day"] for ln in self.linhas()["rad-001"]],
                         ["2026-10-25", "2026-10-26"])

    def test_um_dia_sem_mudanca_grava_na_mesma(self):
        self.sync()
        self.sync(g=guia("2026-10-26", p884113=NORMAL, p898148=SO_FOIL, p1=NORMAL))
        self.assertEqual(len(self.linhas()["rad-113"]), 2)

    def test_o_dia_fica_registado(self):
        self.sync()
        con = self.v.connect()
        r = con.execute("SELECT * FROM prices.cardmarket_dia").fetchall()
        con.close()
        self.assertEqual([(x["day"], x["linhas"], x["origem_ligacao"]) for x in r],
                         [("2026-10-25", 3, "dotgg")])

    def test_nao_toca_na_colecao_nem_no_historico_do_cardtrader(self):
        con = self.v.connect()
        def foto():
            return (con.execute("SELECT * FROM copies ORDER BY printing_id").fetchall(),
                    con.execute("SELECT COUNT(*) FROM ops").fetchone()[0],
                    con.execute("SELECT COUNT(*) FROM prices.price_history").fetchone()[0])
        antes = [list(map(tuple, foto()[0])), foto()[1], foto()[2]]
        self.cmp.sincronizar(con, guia=self.guia, dotgg=self.dotgg, log=lambda *_: None)
        depois = [list(map(tuple, foto()[0])), foto()[1], foto()[2]]
        con.close()
        self.assertEqual(antes, depois)
        self.assertTrue(antes[0], "a fotografia não pode ser de nada")


class TestOPrecoDeReferencia(Base):

    def test_carta_com_normal_usa_o_trend_e_a_foil_o_do_foil(self):
        ln = {"has_normal": 1, "trend_cents": 42, "trend_foil_cents": 212}
        self.assertEqual(self.cmp.preco_referencia(ln), 42)
        self.assertEqual(self.cmp.preco_referencia(ln, foil=True), 212)
        self.assertEqual(self.cmp.valor(ln, 3, 1), 3 * 42 + 212)

    def test_carta_so_foil_usa_o_trend_do_foil_mesmo_nas_normais(self):
        # Ambessa (VEN-196) a 25/10: a página mostra 60,61 € = trend-foil
        ln = {"has_normal": 0, "trend_cents": 9500, "trend_foil_cents": 6061}
        self.assertEqual(self.cmp.preco_referencia(ln), 6061)
        self.assertEqual(self.cmp.valor(ln, 2, 0), 12122)

    def test_sem_trend_o_valor_nao_se_inventa(self):
        ln = {"has_normal": 1, "trend_cents": None, "trend_foil_cents": 212}
        self.assertIsNone(self.cmp.valor(ln, 1, 0))
        self.assertEqual(self.cmp.valor(ln, 0, 1), 212)

    def test_a_variacao_so_com_medicao_dentro_da_janela(self):
        s = [{"day": "2026-10-10", "has_normal": 1, "trend_cents": 10},
             {"day": "2026-10-18", "has_normal": 1, "trend_cents": 30},
             {"day": "2026-10-24", "has_normal": 1, "trend_cents": 40},
             {"day": "2026-10-25", "has_normal": 1, "trend_cents": 50}]
        v = self.cmp.variacao(s)
        self.assertEqual(v["1d"]["cents"], 10)
        self.assertEqual(v["7d"], {"cents": 20, "desde": "2026-10-18", "pct": 66.7})
        # sem a de 18/10, a de 10/10 está fora da folga: não é «7 dias»
        v = self.cmp.variacao([s[0], s[2], s[3]])
        self.assertIsNone(v["7d"])

    def test_o_valor_da_colecao_ao_lado_do_do_site(self):
        self.sync()
        con = self.v.connect()
        v = self.cmp.comparar_valor(con)
        con.close()
        # rad-001: 3 normais × 0,42 + 1 foil × 2,12; rad-113 (só foil): 2 × 60,61
        self.assertEqual(v["cm_cents"], 3 * 42 + 212 + 2 * 6061)
        self.assertEqual(v["copias_sem_trend"], 0)


class TestOExcel(Base):

    def gerar(self):
        from riftvault import excel_precos
        from datetime import datetime
        pasta = ficheiro_temporario(f"precos-{self.id().rsplit('.', 1)[-1]}")
        con = self.v.connect()
        try:
            return excel_precos.gerar(con, pasta, agora=datetime(2026, 10, 25, 17, 15))
        finally:
            con.close()

    def abrir(self, res):
        from openpyxl import load_workbook
        return load_workbook(res["fixo"])

    def test_as_folhas_a_data_e_a_nota(self):
        self.sync()
        res = self.gerar()
        wb = self.abrir(res)
        self.assertEqual(wb.sheetnames, ["Leia-me", "Radiance", "Tudo", "O que tenho",
                                         "A subir - a descer"])
        for nome in wb.sheetnames:
            self.assertIn("Gerado a 2026-10-25 17:15", wb[nome]["A1"].value)
            self.assertIn("preços do Cardmarket de 2026-10-25", wb[nome]["A1"].value)
        self.assertIn("O EXCEL É GERADO, NUNCA EDITADO À MÃO", wb["Leia-me"]["A3"].value)
        self.assertTrue(Path(res["ficheiro"]).name.startswith("riftvault-precos-2026-10-25"))
        self.assertEqual(Path(res["fixo"]).name, "riftvault-precos.xlsx")

    def codigos(self, ws):
        return [r[2] for r in ws.iter_rows(min_row=3, values_only=True) if r[2]]

    def test_a_radiance_a_parte_e_o_que_tenho_so_com_o_que_tem(self):
        self.sync()
        wb = self.abrir(self.gerar())
        self.assertEqual(sorted(self.codigos(wb["Radiance"])), ["RAD-001", "RAD-113"])
        self.assertTrue(all(c.startswith("TST") for c in self.codigos(wb["Tudo"])))
        tenho = [c for c in self.codigos(wb["O que tenho"]) if c != "TOTAL"]
        self.assertEqual(tenho, ["RAD-113", "RAD-001"])   # o mais valioso primeiro

    def test_o_lado_normal_fica_em_branco_numa_carta_so_foil(self):
        self.sync()
        wb = self.abrir(self.gerar())
        ws = wb["Radiance"]
        cab = [c.value for c in ws[2]]
        linha = next(r for r in ws.iter_rows(min_row=3, values_only=True) if r[2] == "RAD-113")
        d = dict(zip(cab, linha))
        self.assertEqual(d["Só foil?"], "sim")
        self.assertIsNone(d["CM Trend €"])
        self.assertIsNone(d["CM média 1d €"])
        self.assertEqual(d["CM Trend foil €"], 60.61)
        self.assertEqual(d["Preço de referência €"], 60.61)
        self.assertEqual(d["Valor das tuas (Trend) €"], 121.22)
        self.assertTrue(d["Link Cardmarket"].startswith("https://www.cardmarket.com/"))

    def test_sem_historico_a_folha_da_variacao_diz_porque(self):
        self.sync()
        wb = self.abrir(self.gerar())
        self.assertIn("Ainda não há variação de 7 dias", wb["A subir - a descer"]["A3"].value)

    def test_gerar_nao_escreve_na_base(self):
        self.sync()
        caminhos = [self.v.data / n for n in ("vault.db", "prices.db", "catalog.db")]
        def h():
            return [hashlib.sha256(p.read_bytes()).hexdigest() for p in caminhos]
        con = self.v.connect(); con.execute("PRAGMA wal_checkpoint(TRUNCATE)"); con.close()
        antes = h()
        self.gerar()
        con = self.v.connect(); con.execute("PRAGMA wal_checkpoint(TRUNCATE)"); con.close()
        self.assertEqual(antes, h())


class TestACli(Base):

    def test_excel_e_valor_sem_rede(self):
        self.sync()
        from riftvault import cli
        pasta = ficheiro_temporario("precos-cli")
        out = io.StringIO()
        with contextlib.redirect_stdout(out):
            self.assertEqual(cli.main(["cardmarket", "--excel", "--pasta", str(pasta)]), 0)
        texto = out.getvalue()
        self.assertIn("riftvault-precos.xlsx", texto)
        self.assertIn("pelo Trend do Cardmarket", texto)
        self.assertIn("mercados diferentes", texto)
        self.assertTrue((pasta / "riftvault-precos.xlsx").exists())


class TestOsSitios(unittest.TestCase):

    def test_a_pasta_do_excel_esta_no_gitignore(self):
        self.assertIn("\nprecos/\n", (REPO / ".gitignore").read_text(encoding="utf-8"))

    def test_o_modulo_diz_o_que_o_cmprice_e(self):
        txt = (REPO / "riftvault" / "cardmarket_precos.py").read_text(encoding="utf-8")
        self.assertIn("NÃO É O TREND — É A MÉDIA DE 1 DIA", txt)

    def test_nenhum_modulo_de_contas_le_os_precos_do_cardmarket(self):
        # é uma vista à parte: nada da Coleção pode passar a depender dela
        for nome in ("metrics", "a_mais", "faltas", "faltas_edicao", "decks", "locais",
                     "prices", "painel", "foil", "collection"):
            txt = (REPO / "riftvault" / f"{nome}.py").read_text(encoding="utf-8")
            self.assertNotIn("cardmarket_precos", txt, nome)
            self.assertNotIn("cardmarket_price_history", txt, nome)


if __name__ == "__main__":
    unittest.main()
