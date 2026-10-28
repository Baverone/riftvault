"""O EXCEL DE PREÇOS — uma VISTA, regenerada todos os dias. Nunca a base de dados.

André, 2026-10-25: *"podes criar Excel para servir de base de dados e ires
atualizando os precos"*.

O EXCEL É GERADO, NUNCA EDITADO À MÃO. A base de dados continua a ser o
`prices.db` e o `vault.db` — se o Excel passasse a ser editável, passavam a
existir duas verdades e elas divergiam na primeira semana (foi o que nos custou
dois dias com o catálogo). O Excel é uma VISTA, regenerada todos os dias e
sempre que ele pedir, com a data e a hora lá dentro. Uma alteração feita no
ficheiro perde-se na geração seguinte; as cópias mudam-se no site.

Cada coluna diz de onde vem e de que dia é. Ver o `cardmarket_precos.py` para
o que cada número do Cardmarket É — em especial: o «Trend» é o `trend` do guia
oficial (o que a página chama «Price Trend»), e numa carta que só existe em
foil é o trend do foil.

Escreve DOIS ficheiros: `riftvault-precos-AAAA-MM-DD.xlsx` (o do dia) e
`riftvault-precos.xlsx` (o mesmo, com nome fixo, para ele abrir sempre o mesmo).
"""

from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path

from . import cardmarket_precos as cmp
from . import config

NOME_FIXO = "riftvault-precos.xlsx"
EDICAO_EM_DESTAQUE = "RAD"

NOTA = (
    "O EXCEL É GERADO, NUNCA EDITADO À MÃO. A base de dados continua a ser o "
    "prices.db e o vault.db — se o Excel passar a ser editável, passam a existir "
    "duas verdades e elas divergem na primeira semana (foi o que nos custou dois "
    "dias com o catálogo). O Excel é uma VISTA, regenerada todos os dias e sempre "
    "que pedires, com a data e a hora lá dentro. O que mudares aqui perde-se na "
    "geração seguinte: as cópias mudam-se no site.")

FONTES = [
    ("CM Trend", "o «Price Trend» da página da carta no Cardmarket — o campo `trend` "
                 "do price guide oficial do Cardmarket (ficheiro de download diário, "
                 "jogo 22). Numa carta que só existe em foil, a página mostra o do foil."),
    ("CM From", "a oferta mais barata (`low`)."),
    ("CM média 1/7/30 dias", "as três médias de venda da página (`avg1`, `avg7`, `avg30`)."),
    ("Preço de referência", "o que a página mostra como Trend: o normal se a carta tem "
                            "versão normal; o do foil se só existe em foil."),
    ("Var. Trend 1d / 7d", "a variação do preço de referência contra o que o riftvault "
                           "guardou há 1 e há 7 dias (com 1–2 dias de folga). «—» = ainda "
                           "não há medição desse dia. Começou a 2026-10-25."),
    ("CardTrader", "a oferta mais barata em inglês, NM/Mint, do CardTrader — o preço que o "
                   "site usa. Outro mercado: não se compara um com o outro como se fossem o mesmo."),
    ("Outro mercado (USD)", "o `price` da dotgg, em DÓLARES; não se converte."),
    ("NÃO É o cmPrice da dotgg", "o `cmPrice` da dotgg é a MÉDIA DE 1 DIA do Cardmarket "
                                 "(medido a 2026-10-25: 1429 de 1633 batem ao cêntimo com o "
                                 "avg1), não o Trend. Não se usa."),
]

COLUNAS_CARTA = [
    ("Edição", 8), ("Nº", 6), ("Código", 14), ("Nome", 32), ("Versão", 13),
    ("Raridade", 11), ("Tipo", 11), ("Tens normais", 8), ("Tens foil", 8),
    ("Só foil?", 7), ("Preço de referência €", 12), ("Var. 1d €", 9), ("Var. 7d €", 9),
    ("Var. 7d %", 9), ("CM Trend €", 10), ("CM Trend foil €", 10),
    ("CM From €", 9), ("CM From foil €", 9), ("CM média 1d €", 10),
    ("CM média 7d €", 10), ("CM média 30d €", 10), ("CM média 1d foil €", 10),
    ("CM média 7d foil €", 10), ("CM média 30d foil €", 10),
    ("CardTrader €", 10), ("CardTrader foil €", 10), ("Outro mercado USD", 10),
    ("Outro mercado foil USD", 10), ("Valor das tuas (Trend) €", 12),
    ("Dia do preço CM", 11), ("Link Cardmarket", 50),
]


def _linhas(con) -> list[dict]:
    """Uma linha por impressão do catálogo, com o último preço do Cardmarket."""
    series = cmp.series(con)
    copias = {r["printing_id"]: (r["qty"] or 0, r["qty_foil"] or 0)
              for r in con.execute("SELECT printing_id, qty, qty_foil FROM copies")}
    ct = {r["printing_id"]: (r["price_cents"], r["price_foil_cents"])
          for r in con.execute("SELECT printing_id, price_cents, price_foil_cents "
                               "FROM catalog.price_latest")}
    out = []
    impressoes = [dict(p) for p in con.execute(
        "SELECT printing_id, set_id, collector_number, public_code, name, "
        "variant_label, base_rarity, rarity, type, api_sort FROM catalog.printings")]
    # As do CardTrader que a RiftScribe não tem (`catalog.market_only`, as runas
    # em alt art do SFD/UNL/VEN): ele TEM algumas (SFD-R02a ×12), e uma folha
    # «O que tenho» sem elas mentia por omissão. Vão para o fim da edição. SÓ AS
    # QUE ELE TEM: a tabela traz também restos de mapeamentos antigos que são a
    # mesma carta de uma impressão do catálogo (as SP da RAD, medido a 25/10), e
    # listá-las todas punha linhas repetidas na folha.
    for m in con.execute("SELECT m.printing_id, m.set_id, m.collector_raw, m.market_name "
                         "FROM catalog.market_only m JOIN copies c "
                         "  ON c.printing_id = m.printing_id "
                         "WHERE c.qty > 0 OR c.qty_foil > 0"):
        sid = (m["set_id"] or "").upper()
        impressoes.append({"printing_id": m["printing_id"], "set_id": sid,
                           "collector_number": None,
                           "public_code": f"{sid}-{m['collector_raw']}".upper(),
                           "name": m["market_name"], "variant_label": "fora do catálogo",
                           "base_rarity": None, "rarity": None, "type": None,
                           "api_sort": 100000})
    for p in impressoes:
        serie = series.get(p["printing_id"]) or []
        ln = serie[-1] if serie else {}
        qn, qf = copias.get(p["printing_id"], (0, 0))
        var = cmp.variacao(serie)
        out.append({"p": p, "cm": ln, "qn": qn, "qf": qf, "var": var,
                    "ct": ct.get(p["printing_id"], (None, None)),
                    "valor": cmp.valor(ln, qn, qf) if ln and (qn or qf) else None})
    out.sort(key=lambda r: (config.set_order(r["p"]["set_id"]), r["p"]["set_id"],
                            r["p"]["api_sort"] if r["p"]["api_sort"] is not None else 0))
    return out


def _celulas(r: dict) -> list:
    p, ln, var = r["p"], r["cm"], r["var"]
    e = cmp.eur
    # Numa carta que só existe em foil, o lado NORMAL do guia é ruído (trend 0,
    # um «From» que é o do foil, às vezes uma média solta — OGN-119 tinha média
    # de 1 dia 8,00 € no lado normal e a página mostra 4,58 €). Fica em branco
    # em vez de parecer um preço.
    so_foil = bool(ln) and not ln.get("has_normal")
    n = (lambda k: None) if so_foil else (lambda k: e(ln.get(k)))
    v7 = var.get("7d") or {}
    url = ln.get("cm_url")
    if not url and ln:
        from . import mercados
        url = mercados.cardmarket_busca(p["name"])
    return [
        p["set_id"], p["collector_number"], p["public_code"], p["name"], p["variant_label"],
        p["base_rarity"] or p["rarity"], p["type"], r["qn"], r["qf"],
        ("" if not ln else ("não" if ln.get("has_normal") else "sim")),
        e(cmp.preco_referencia(ln)) if ln else None,
        e((var.get("1d") or {}).get("cents")), e(v7.get("cents")), v7.get("pct"),
        n("trend_cents"), e(ln.get("trend_foil_cents")),
        n("low_cents"), e(ln.get("low_foil_cents")),
        n("avg1_cents"), n("avg7_cents"), n("avg30_cents"),
        e(ln.get("avg1_foil_cents")), e(ln.get("avg7_foil_cents")), e(ln.get("avg30_foil_cents")),
        e(r["ct"][0]), e(r["ct"][1]),
        e(ln.get("outro_usd_cents")), e(ln.get("outro_usd_foil_cents")),
        e(r["valor"]), ln.get("day"), url,
    ]


def gerar(con, pasta: Path | None = None, agora: datetime | None = None) -> dict:
    """Gera os dois ficheiros. Só LÊ a base — não escreve em sítio nenhum dela."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    pasta = Path(pasta or config.PRECOS_DIR)
    pasta.mkdir(parents=True, exist_ok=True)
    agora = agora or datetime.now().replace(microsecond=0)
    linhas = _linhas(con)
    dias = sorted({r["cm"].get("day") for r in linhas if r["cm"].get("day")})
    dia_cm = dias[-1] if dias else None
    titulo = (f"Gerado a {agora:%Y-%m-%d %H:%M} · preços do Cardmarket de "
              f"{dia_cm or '— (ainda não há)'} · NÃO EDITAR: regenera-se")

    wb = Workbook()
    negrito = Font(bold=True)
    cab = PatternFill("solid", fgColor="2B2140")
    cab_f = Font(bold=True, color="FFFFFF")

    def folha(nome, filas, cols=COLUNAS_CARTA):
        ws = wb.create_sheet(nome)
        ws.append([titulo])
        ws["A1"].font = negrito
        ws.append([c for c, _ in cols])
        for i, (_, larg) in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(i)].width = larg
            c = ws.cell(row=2, column=i)
            c.font, c.fill = cab_f, cab
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for f in filas:
            ws.append(f)
        ws.freeze_panes = "E3"
        ws.auto_filter.ref = f"A2:{get_column_letter(len(cols))}{max(2, ws.max_row)}"
        if cols is COLUNAS_CARTA:
            ilink = len(cols)
            for row in ws.iter_rows(min_row=3, min_col=ilink, max_col=ilink):
                for c in row:
                    if c.value:
                        c.hyperlink = c.value
                        c.style = "Hyperlink"
        return ws

    # Leia-me primeiro: é a primeira coisa que ele vê ao abrir.
    ws = wb.active
    ws.title = "Leia-me"
    ws.append([titulo]); ws["A1"].font = negrito
    ws.append([])
    ws.append([NOTA]); ws["A3"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[3].height = 90
    ws.append([])
    ws.append(["De onde vem cada coluna"]); ws["A5"].font = negrito
    for nome, txt in FONTES:
        ws.append([nome, txt])
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 120
    for row in ws.iter_rows(min_row=6, max_col=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    rad = [r for r in linhas if r["p"]["set_id"] == EDICAO_EM_DESTAQUE]
    outras = [r for r in linhas if r["p"]["set_id"] != EDICAO_EM_DESTAQUE]
    folha("Radiance", [_celulas(r) for r in rad])
    folha("Tudo", [_celulas(r) for r in outras])

    tenho = [r for r in linhas if r["qn"] or r["qf"]]
    tenho.sort(key=lambda r: -(r["valor"] or 0))
    ws_t = folha("O que tenho", [_celulas(r) for r in tenho])
    total = sum(r["valor"] or 0 for r in tenho)
    sem = sum(1 for r in tenho if r["valor"] is None)
    ws_t.append([])
    ws_t.append(["TOTAL", None, None,
                 f"{cmp.eur(total):.2f} € pelo Trend do Cardmarket, {len(tenho)} impressões, "
                 f"TODAS as cópias que tens (também as dos decks — o valor do site não "
                 f"conta as próprias dos decks nem o stock)"
                 + (f" — {sem} sem Trend ficam de fora (não se inventa)" if sem else "")])
    ws_t.cell(row=ws_t.max_row, column=1).font = negrito

    # A subir / a descer: pela variação de 7 dias do preço de referência.
    com_var = [r for r in linhas if (r["var"].get("7d") or {}).get("cents") is not None]
    com_var.sort(key=lambda r: -r["var"]["7d"]["cents"])
    ws_v = folha("A subir - a descer", [_celulas(r) for r in com_var])
    if not com_var:
        ws_v.append(["Ainda não há variação de 7 dias: o riftvault começou a guardar os "
                     "preços do Cardmarket a 2026-10-25. A primeira aparece uma semana "
                     "depois — não se pede a variação a outra fonte para encher a folha."])

    nome_dia = f"riftvault-precos-{agora:%Y-%m-%d}.xlsx"
    caminho = pasta / nome_dia
    wb.save(caminho)
    fixo = pasta / NOME_FIXO
    fixo_falhou = None
    try:
        shutil.copyfile(caminho, fixo)
    except PermissionError as exc:
        # O Windows tranca um .xlsx aberto no Excel: fica o do dia, e diz-se.
        fixo_falhou = f"{fixo.name} está aberto noutro programa — ficou só o do dia ({exc})"
    return {"ficheiro": str(caminho), "fixo": str(fixo), "fixo_falhou": fixo_falhou,
            "dia_cm": dia_cm,
            "linhas": len(linhas), "rad": len(rad), "tenho": len(tenho),
            "valor_cents": total, "sem_trend_tenho": sem, "a_subir": len(com_var)}
